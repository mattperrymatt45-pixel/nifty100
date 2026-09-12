"""Tests for the Day 26 sector-relative valuation module."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest

from src.analytics import valuation as val

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DB_PATH = PROJECT_ROOT / "db" / "nifty100.db"


@pytest.fixture(scope="module")
def panel() -> pd.DataFrame:
    return val.load_valuation_panel(str(DB_PATH))


@pytest.fixture(scope="module")
def summary(panel: pd.DataFrame) -> pd.DataFrame:
    return val.compute_valuation_summary(panel)


# ---------------------------------------------------------------------------
# Primitive sanity (FCF yield formula)
# ---------------------------------------------------------------------------
def test_fcf_yield_positive() -> None:
    assert val.fcf_yield(100.0, 2000.0) == pytest.approx(5.0)


def test_fcf_yield_zero_market_cap_returns_none() -> None:
    assert val.fcf_yield(100.0, 0) is None


def test_fcf_yield_negative_fcf() -> None:
    assert val.fcf_yield(-50.0, 1000.0) == pytest.approx(-5.0)


# ---------------------------------------------------------------------------
# Panel shape
# ---------------------------------------------------------------------------
def test_panel_covers_all_latest_companies(panel: pd.DataFrame) -> None:
    assert len(panel) == 89
    for col in [
        "company_id",
        "company_name",
        "sector",
        "pe_ratio",
        "pb_ratio",
        "ev_ebitda",
        "free_cash_flow_cr",
        "market_cap_crore",
        "pe_5yr_median",
    ]:
        assert col in panel.columns


# ---------------------------------------------------------------------------
# Summary invariants
# ---------------------------------------------------------------------------
def test_summary_columns_match_spec(summary: pd.DataFrame) -> None:
    assert tuple(summary.columns) == val.VALUATION_SUMMARY_COLUMNS


def test_summary_fcf_yield_formula(summary: pd.DataFrame, panel: pd.DataFrame) -> None:
    """FCF Yield = FCF / mcap * 100 (spot check)."""
    merged = summary.merge(
        panel[["company_id", "free_cash_flow_cr", "market_cap_crore"]], on="company_id"
    )
    sample = merged.dropna(subset=["free_cash_flow_cr", "market_cap_crore"]).iloc[0]
    expected = sample["free_cash_flow_cr"] / sample["market_cap_crore"] * 100
    assert abs(sample["fcf_yield_pct"] - expected) < 0.01


def test_summary_flag_distribution(summary: pd.DataFrame) -> None:
    flags = set(summary["flag"].unique())
    assert flags <= {"Caution", "Discount", "Fair"}
    # We expect a reasonable mix given the Nifty 100 universe.
    counts = summary["flag"].value_counts()
    assert counts.get("Fair", 0) >= 20
    assert counts.get("Caution", 0) >= 5
    assert counts.get("Discount", 0) >= 5


def test_sector_median_caution_flag_logic() -> None:
    """Synthetic sector: median=20 → PE>30 must be Caution, PE<14 Discount, else Fair."""
    df = pd.DataFrame(
        {
            "company_id": ["A", "B", "C"],
            "company_name": ["A Ltd", "B Ltd", "C Ltd"],
            "sector": ["Tech"] * 3,
            "pe_ratio": [35.0, 10.0, 20.0],
            "pb_ratio": [5.0, 1.0, 2.5],
            "ev_ebitda": [20.0, 8.0, 12.0],
            "free_cash_flow_cr": [100.0, 50.0, 80.0],
            "market_cap_crore": [1000.0, 500.0, 800.0],
            "pe_5yr_median": [30.0, 12.0, 18.0],
        }
    )
    out = val.compute_valuation_summary(df)
    assert out.loc[out["company_id"] == "A", "flag"].iloc[0] == "Caution"
    assert out.loc[out["company_id"] == "B", "flag"].iloc[0] == "Discount"
    assert out.loc[out["company_id"] == "C", "flag"].iloc[0] == "Fair"


def test_negative_pe_defaults_to_fair(panel: pd.DataFrame) -> None:
    """Loss-makers (PE NaN / <= 0) should be flagged Fair — insufficient data."""
    df = panel.copy()
    df.loc[df.index[0], "pe_ratio"] = -5.0
    out = val.compute_valuation_summary(df)
    assert out.iloc[0]["flag"] == "Fair"


def test_fcf_yield_pct_stored(summary: pd.DataFrame) -> None:
    """fcf_yield_pct column must be present and non-null for most companies."""
    non_null = summary["fcf_yield_pct"].notna().sum()
    assert non_null >= 80


def test_5yr_median_pe_present(summary: pd.DataFrame) -> None:
    """5yr_median_PE should be populated for the vast majority."""
    assert summary["5yr_median_PE"].notna().sum() >= 80


# ---------------------------------------------------------------------------
# Output files
# ---------------------------------------------------------------------------
def test_write_outputs_creates_xlsx_and_csv(tmp_path: Path, summary: pd.DataFrame) -> None:
    xlsx, csv = val.write_valuation_outputs(summary, tmp_path)
    assert xlsx.exists() and xlsx.stat().st_size > 2000
    assert csv.exists() and csv.stat().st_size > 50

    # CSV must contain only Caution/Discount, no Fair
    flagged = pd.read_csv(csv)
    assert set(flagged["flag"]) <= {"Caution", "Discount"}
    assert len(flagged) >= 20

    # XLSX should be readable openpyxl
    from openpyxl import load_workbook

    wb = load_workbook(xlsx)
    assert "Valuation Summary" in wb.sheetnames
    ws = wb["Valuation Summary"]
    assert ws.cell(row=1, column=1).value == "company_id"
    # Header row should be 11 cells wide (VALUATION_SUMMARY_COLUMNS count)
    headers = [ws.cell(row=1, column=i).value for i in range(1, 12)]
    assert headers == list(val.VALUATION_SUMMARY_COLUMNS)


def test_end_to_end_run(tmp_path: Path) -> None:
    summary, flagged, xlsx, csv = val.run_valuation_module(db_path=DB_PATH, output_dir=tmp_path)
    assert len(summary) == 89
    assert len(flagged) >= 20
    assert xlsx.exists()
    assert csv.exists()
