"""Sprint 5 Day 31 — Cash Flow Intelligence Module.

Extends the Day-11 Cash Flow KPI primitives in :mod:`src.analytics.cashflow_kpis`
with a full panel builder and Excel/CSV writers that produce:

* ``output/cashflow_intelligence.xlsx`` — one row per Nifty 100 company with
  CFO quality score/label, CapEx intensity/label, 5yr FCF CAGR, FCF conversion
  rate, distress flag, deleveraging flag, and capital-allocation pattern label.
* ``output/distress_alerts.csv`` — companies flagged as distressed in the
  latest year (CFO < 0 AND CFF > 0), with CFO/CFF values and latest net profit.
"""

from __future__ import annotations

import sqlite3
from dataclasses import dataclass
from pathlib import Path

import pandas as pd

from src.analytics.cashflow_kpis import (
    capex_intensity,
    capex_tier,
    cfo_pat_ratio,
    cfo_quality_tier,
    classify_capital_allocation,
    fcf_conversion,
    free_cash_flow,
)

# ---------------------------------------------------------------------------
# Output schemas
# ---------------------------------------------------------------------------
CASHOUTPUT_COLUMNS: tuple[str, ...] = (
    "company_id",
    "company_name",
    "sector",
    "cfo_quality_score",
    "cfo_quality_label",
    "capex_intensity_pct",
    "capex_label",
    "fcf_cagr_5yr",
    "fcf_conversion_pct",
    "distress_flag",
    "deleveraging_flag",
    "capital_allocation_label",
)

DISTRESS_COLUMNS: tuple[str, ...] = (
    "company_id",
    "company_name",
    "sector",
    "year",
    "cfo_cr",
    "cff_cr",
    "net_profit_cr",
)


@dataclass(frozen=True)
class DistressAlertRow:
    """One row of output/distress_alerts.csv."""

    company_id: str
    company_name: str
    sector: str | None
    year: str
    cfo_cr: float
    cff_cr: float
    net_profit_cr: float

    def as_dict(self) -> dict:
        return {
            "company_id": self.company_id,
            "company_name": self.company_name,
            "sector": self.sector,
            "year": self.year,
            "cfo_cr": round(self.cfo_cr, 2),
            "cff_cr": round(self.cff_cr, 2),
            "net_profit_cr": round(self.net_profit_cr, 2),
        }


# ---------------------------------------------------------------------------
# New primitive flags / calculations
# ---------------------------------------------------------------------------
def distress_signal(cfo: float | None, cff: float | None) -> bool:
    """Day-31 distress flag: CFO < 0 AND CFF > 0 in the latest year.

    Returns True only when both flows are known and satisfy the predicate;
    missing data returns False so we don't raise false positives.
    """
    if cfo is None or cff is None:
        return False
    try:
        return float(cfo) < 0 and float(cff) > 0
    except (TypeError, ValueError):
        return False


def deleveraging_flag(
    cff: float | None,
    borrowings_now: float | None,
    borrowings_prev: float | None,
) -> bool:
    """Day-31 deleveraging flag: CFF < 0 AND borrowings declined YoY.

    CFF negative means the company repaid debt / bought back stock / paid
    dividends; combined with a genuine drop in gross borrowings this is a
    clean deleveraging signal.  Missing data returns False.
    """
    if cff is None:
        return False
    try:
        cff_f = float(cff)
    except (TypeError, ValueError):
        return False
    if cff_f >= 0:
        return False
    if borrowings_now is None or borrowings_prev is None:
        return False
    try:
        b_now = float(borrowings_now)
        b_prev = float(borrowings_prev)
    except (TypeError, ValueError):
        return False
    if pd.isna(b_now) or pd.isna(b_prev):
        return False
    return b_now < b_prev


def fcf_cagr(fcf_series: list[float | None], window: int = 5) -> float | None:
    """Compound annual growth rate of FCF over the trailing ``window`` years.

    Uses (end/begin)**(1/n) - 1 on FCF values.  Requires begin- and end-year
    FCF strictly positive (negative FCF makes CAGR mathematically undefined).
    """
    vals = [v for v in fcf_series if v is not None and not pd.isna(v)]
    # Take up to window+1 points for a window-year CAGR (need end vs begin)
    vals = vals[-(window + 1) :]
    if len(vals) < 2:
        return None
    begin = float(vals[0])
    end = float(vals[-1])
    n = len(vals) - 1
    if begin <= 0 or end <= 0 or n <= 0:
        return None
    try:
        return ((end / begin) ** (1.0 / n) - 1.0) * 100.0
    except (ValueError, ZeroDivisionError):
        return None


# ---------------------------------------------------------------------------
# Panel builder
# ---------------------------------------------------------------------------
_QUERY = """
    SELECT
        c.id AS company_id,
        c.company_name,
        s.broad_sector AS sector,
        cf.year,
        cf.operating_activity   AS cfo,
        cf.investing_activity   AS cfi,
        cf.financing_activity   AS cff,
        pl.sales,
        pl.operating_profit,
        pl.net_profit,
        bs.borrowings
    FROM companies c
    JOIN cashflow cf ON cf.company_id = c.id
    LEFT JOIN profitandloss pl ON pl.company_id = c.id AND pl.year = cf.year
    LEFT JOIN balancesheet bs ON bs.company_id = c.id AND bs.year = cf.year
    LEFT JOIN sectors s ON s.company_id = c.id
    ORDER BY c.company_name, cf.year
"""


def build_cashflow_intelligence_panel(
    conn: sqlite3.Connection | str | Path,
) -> pd.DataFrame:
    """Return the cashflow-intelligence summary (one row per company, latest FY)."""
    if isinstance(conn, (str, Path)):
        conn = sqlite3.connect(str(conn))
        close_conn = True
    else:
        close_conn = False
    try:
        df = pd.read_sql(_QUERY, conn)
    finally:
        if close_conn:
            conn.close()

    rows: list[dict] = []
    for cid, g in df.groupby("company_id", sort=False):
        g = g.sort_values("year").reset_index(drop=True)
        company_name = g["company_name"].iloc[0]
        sector = g["sector"].iloc[0]
        if isinstance(sector, float) and pd.isna(sector):
            sector = None

        cfo_pat_hist: list[float | None] = []
        fcf_hist: list[float | None] = []
        borrowings_list: list = []

        for _, r in g.iterrows():
            cfo = r["cfo"]
            cfi = r["cfi"]
            cff = r["cff"]
            pat = r["net_profit"]
            borrowings = r["borrowings"]
            borrowings_list.append(borrowings)

            if any(pd.isna(x) for x in (cfo, cfi, cff)):
                cfo_pat_hist.append(None)
                fcf_hist.append(None)
                continue
            cfo_f, cfi_f = float(cfo), float(cfi)
            fcf_f = free_cash_flow(cfo_f, cfi_f)
            if pd.isna(pat):
                cfo_pat_hist.append(None)
            else:
                cfo_pat_hist.append(cfo_pat_ratio(cfo_f, float(pat)))
            fcf_hist.append(fcf_f)

        latest_idx = len(g) - 1
        cfo_latest = g.at[latest_idx, "cfo"]
        cfi_latest = g.at[latest_idx, "cfi"]
        cff_latest = g.at[latest_idx, "cff"]
        sales_latest = g.at[latest_idx, "sales"]
        op_latest = g.at[latest_idx, "operating_profit"]
        pat_latest = g.at[latest_idx, "net_profit"]

        cfo_quality_score: float | None = None
        cfo_quality_label: str | None = None
        capex_pct: float | None = None
        capex_lbl: str | None = None
        fcf_conv: float | None = None
        pattern_label: str | None = None

        if not any(pd.isna(x) for x in (cfo_latest, cfi_latest, cff_latest)):
            cfo_f, cfi_f, cff_f = float(cfo_latest), float(cfi_latest), float(cff_latest)
            fcf_latest = free_cash_flow(cfo_f, cfi_f)
            cp = cfo_pat_ratio(cfo_f, float(pat_latest)) if not pd.isna(pat_latest) else None
            _s1, _s2, _s3, pattern_label = classify_capital_allocation(cfo_f, cfi_f, cff_f, cp)

            window_vals = [v for v in cfo_pat_hist[-5:] if v is not None and not pd.isna(v)]
            if len(window_vals) >= 3:
                cfo_quality_score = sum(window_vals) / len(window_vals)
                cfo_quality_label = cfo_quality_tier(cfo_quality_score)
            if not pd.isna(sales_latest) and float(sales_latest) != 0:
                capex_pct = capex_intensity(cfi_f, float(sales_latest))
                capex_lbl = capex_tier(capex_pct)
            if not pd.isna(op_latest) and float(op_latest) != 0:
                fcf_conv = fcf_conversion(fcf_latest, float(op_latest))

        fcf_5yr_cagr = fcf_cagr(fcf_hist, window=5)
        distress = distress_signal(cfo_latest, cff_latest)
        b_now = borrowings_list[-1] if len(borrowings_list) >= 1 else None
        b_prev = borrowings_list[-2] if len(borrowings_list) >= 2 else None
        deleveraging = deleveraging_flag(cff_latest, b_now, b_prev)

        rows.append(
            {
                "company_id": cid,
                "company_name": company_name,
                "sector": sector,
                "cfo_quality_score": (
                    round(cfo_quality_score, 2) if cfo_quality_score is not None else None
                ),
                "cfo_quality_label": cfo_quality_label,
                "capex_intensity_pct": round(capex_pct, 2) if capex_pct is not None else None,
                "capex_label": capex_lbl,
                "fcf_cagr_5yr": round(fcf_5yr_cagr, 2) if fcf_5yr_cagr is not None else None,
                "fcf_conversion_pct": round(fcf_conv, 2) if fcf_conv is not None else None,
                "distress_flag": distress,
                "deleveraging_flag": deleveraging,
                "capital_allocation_label": pattern_label,
            }
        )

    return pd.DataFrame(rows, columns=list(CASHOUTPUT_COLUMNS))


# ---------------------------------------------------------------------------
# Writers
# ---------------------------------------------------------------------------
def write_intelligence_xlsx(df: pd.DataFrame, path: Path) -> Path:
    """Write cashflow-intelligence summary to XLSX with navy header + flag fills."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils.dataframe import dataframe_to_rows

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Cashflow Intelligence"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    true_fill = PatternFill("solid", fgColor="F4CCCC")
    false_fill = PatternFill("solid", fgColor="D9EAD3")

    cols = list(CASHOUTPUT_COLUMNS)
    out_df = df[cols]
    for r_idx, row in enumerate(dataframe_to_rows(out_df, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            if r_idx == 1:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")
            else:
                col_name = cols[c_idx - 1]
                if col_name in ("distress_flag", "deleveraging_flag"):
                    cell.fill = true_fill if bool(value) else false_fill

    for col_cells in ws.columns:
        max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 3, 30)

    ws.freeze_panes = "A2"
    wb.save(path)
    return path


def write_distress_alerts_csv(df: pd.DataFrame, conn: sqlite3.Connection, path: Path) -> int:
    """Write distress-alert rows (companies flagged with distress_signal) to CSV."""
    path.parent.mkdir(parents=True, exist_ok=True)
    flagged = df[df["distress_flag"] == True].copy()  # noqa: E712
    alerts: list[dict] = []
    cur = conn.cursor()
    for _, r in flagged.iterrows():
        cid = r["company_id"]
        row = cur.execute(
            """
            SELECT cf.year, cf.operating_activity, cf.financing_activity, pl.net_profit
            FROM cashflow cf
            LEFT JOIN profitandloss pl ON pl.company_id = cf.company_id AND pl.year = cf.year
            WHERE cf.company_id = ?
            ORDER BY cf.year DESC LIMIT 1
            """,
            (cid,),
        ).fetchone()
        if row is None:
            continue
        yr, cfo_cr, cff_cr, net_profit_cr = row
        alerts.append(
            DistressAlertRow(
                company_id=cid,
                company_name=r["company_name"],
                sector=r["sector"],
                year=yr,
                cfo_cr=float(cfo_cr) if cfo_cr is not None else float("nan"),
                cff_cr=float(cff_cr) if cff_cr is not None else float("nan"),
                net_profit_cr=float(net_profit_cr) if net_profit_cr is not None else float("nan"),
            ).as_dict()
        )
    out_df = pd.DataFrame(alerts, columns=list(DISTRESS_COLUMNS))
    out_df.to_csv(path, index=False)
    return len(out_df)


# ---------------------------------------------------------------------------
# High-level entry point
# ---------------------------------------------------------------------------
def run_cashflow_intelligence(
    db_path: Path | str | None = None,
    output_dir: Path | str | None = None,
) -> tuple[pd.DataFrame, pd.DataFrame, Path, Path]:
    """Run the Day-31 cash flow intelligence module.

    Returns ``(summary_df, alerts_df, xlsx_path, csv_path)``.
    """
    from src.utils.config import settings

    if db_path is None:
        db_path = settings.PROJECT_ROOT / "db" / "nifty100.db"
    if output_dir is None:
        output_dir = settings.PROJECT_ROOT / "output"
    db_path = Path(db_path)
    output_dir = Path(output_dir)

    conn = sqlite3.connect(str(db_path))
    try:
        summary = build_cashflow_intelligence_panel(conn)
        xlsx_path = write_intelligence_xlsx(summary, output_dir / "cashflow_intelligence.xlsx")
        write_distress_alerts_csv(summary, conn, output_dir / "distress_alerts.csv")
        alerts = pd.read_csv(output_dir / "distress_alerts.csv")
    finally:
        conn.close()
    return summary, alerts, xlsx_path, output_dir / "distress_alerts.csv"
