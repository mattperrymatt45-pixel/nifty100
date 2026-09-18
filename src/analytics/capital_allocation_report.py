"""Sprint 5 Day 32 — Capital Allocation Report Module.

Builds on the Day-11 capital-allocation CSV (output/capital_allocation.csv) and
the Day-31 cashflow-intelligence panel to produce:

* **Completeness verification** — confirm capital_allocation.csv covers every
  Nifty-100 company x every shared cashflow+P&L year.
* **Distribution summary** — count of companies in each of the 8 capital-allocation
  patterns for the latest available FY.
* **Pattern-changes CSV** (``output/pattern_changes.csv``) — companies whose
  capital-allocation pattern shifted year-over-year, with before/after labels.
* **Cashflow-Intelligence XLSX refresh** — ensures the ``capital_allocation_label``
  column is present and populated for every company.

The 8-class capital-allocation taxonomy (per spec §13 / Day 11):

    +----+----+----+-----------------------------------------------+
    |CFO |CFI |CFF | Pattern label                                 |
    +====+====+====+===============================================+
    | +  | -  | -  | Reinvestor OR Shareholder Returns (CFO/PAT≥1) |
    +----+----+----+-----------------------------------------------+
    | +  | +  | -  | Liquidating Assets                            |
    +----+----+----+-----------------------------------------------+
    | -  | +  | +  | Distress Signal                               |
    +----+----+----+-----------------------------------------------+
    | -  | -  | +  | Growth Funded by Debt                         |
    +----+----+----+-----------------------------------------------+
    | +  | +  | +  | Cash Accumulator                              |
    +----+----+----+-----------------------------------------------+
    | -  | -  | -  | Pre-Revenue                                   |
    +----+----+----+-----------------------------------------------+
    |other combos|  | Mixed                                         |
    +----+----+----+-----------------------------------------------+
"""

from __future__ import annotations

import logging
import sqlite3
from dataclasses import dataclass
from pathlib import Path

import pandas as pd

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Constants — 8 canonical patterns (mirrors src.analytics.cashflow_kpis)
# ---------------------------------------------------------------------------
PATTERN_REINVESTOR = "Reinvestor"
PATTERN_SHAREHOLDER_RETURNS = "Shareholder Returns"
PATTERN_LIQUIDATING_ASSETS = "Liquidating Assets"
PATTERN_DISTRESS_SIGNAL = "Distress Signal"
PATTERN_GROWTH_FUNDED_BY_DEBT = "Growth Funded by Debt"
PATTERN_CASH_ACCUMULATOR = "Cash Accumulator"
PATTERN_PRE_REVENUE = "Pre-Revenue"
PATTERN_MIXED = "Mixed"

ALL_EIGHT_PATTERNS: tuple[str, ...] = (
    PATTERN_SHAREHOLDER_RETURNS,
    PATTERN_REINVESTOR,
    PATTERN_GROWTH_FUNDED_BY_DEBT,
    PATTERN_DISTRESS_SIGNAL,
    PATTERN_LIQUIDATING_ASSETS,
    PATTERN_CASH_ACCUMULATOR,
    PATTERN_PRE_REVENUE,
    PATTERN_MIXED,
)

PATTERN_CHANGES_COLUMNS: tuple[str, ...] = (
    "company_id",
    "company_name",
    "sector",
    "prev_year",
    "prev_pattern",
    "latest_year",
    "latest_pattern",
)


@dataclass(frozen=True)
class PatternChangeRow:
    """One row in output/pattern_changes.csv (a YoY pattern transition)."""

    company_id: str
    company_name: str
    sector: str | None
    prev_year: str
    prev_pattern: str
    latest_year: str
    latest_pattern: str

    def as_dict(self) -> dict[str, str | None]:
        """Return the row as a flat dictionary suitable for CSV writing."""
        return {
            "company_id": self.company_id,
            "company_name": self.company_name,
            "sector": self.sector,
            "prev_year": self.prev_year,
            "prev_pattern": self.prev_pattern,
            "latest_year": self.latest_year,
            "latest_pattern": self.latest_pattern,
        }


# ---------------------------------------------------------------------------
# Completeness verification
# ---------------------------------------------------------------------------
def verify_capital_allocation_completeness(
    ca_df: pd.DataFrame,
    conn: sqlite3.Connection,
) -> dict:
    """Verify ``capital_allocation.csv`` covers every company x shared year.

    Returns a dict with keys:
        * ``total_companies_db``
        * ``total_companies_csv``
        * ``total_rows_expected``
        * ``total_rows_csv``
        * ``missing_company_ids`` (list of company_ids absent from CSV)
        * ``missing_rows`` (list of (company_id, year) tuples)
        * ``is_complete`` (bool)
    """
    companies_df = pd.read_sql("SELECT id FROM companies ORDER BY id", conn)
    all_cids = set(companies_df["id"].tolist())
    csv_cids = set(ca_df["company_id"].unique().tolist())

    # Expected rows = (company, year) pairs present in BOTH cashflow and P&L
    expected_df = pd.read_sql(
        """
        SELECT cf.company_id, cf.year
        FROM cashflow cf
        INNER JOIN profitandloss pl
            ON pl.company_id = cf.company_id AND pl.year = cf.year
        """,
        conn,
    )
    expected_pairs = set(zip(expected_df["company_id"], expected_df["year"], strict=True))
    csv_pairs = set(zip(ca_df["company_id"], ca_df["year"], strict=True))

    missing_companies = sorted(all_cids - csv_cids)
    missing_rows = sorted(expected_pairs - csv_pairs)

    result = {
        "total_companies_db": len(all_cids),
        "total_companies_csv": len(csv_cids),
        "total_rows_expected": len(expected_pairs),
        "total_rows_csv": len(csv_pairs),
        "missing_company_ids": missing_companies,
        "missing_rows": missing_rows,
        "is_complete": (not missing_companies) and (not missing_rows),
    }
    logger.info(
        "Capital allocation CSV completeness: %s companies, %s/%s rows — is_complete=%s",
        result["total_companies_csv"],
        result["total_rows_csv"],
        result["total_rows_expected"],
        result["is_complete"],
    )
    return result


# ---------------------------------------------------------------------------
# Distribution summary
# ---------------------------------------------------------------------------
def build_pattern_distribution(ca_df: pd.DataFrame) -> pd.DataFrame:
    """Return a DataFrame with count of companies in each pattern for latest FY.

    The "latest FY" per company is that company's maximum year (handles late
    filers).  Companies with NO latest-year pattern (shouldn't happen but safe)
    are excluded.  All 8 canonical patterns appear in the result even if count=0.
    """
    # Latest year per company
    idx = ca_df.groupby("company_id")["year"].idxmax()
    latest = ca_df.loc[idx, ["company_id", "year", "pattern_label"]].copy()

    counts = latest["pattern_label"].value_counts().to_dict()
    rows: list[dict] = []
    for pattern in ALL_EIGHT_PATTERNS:
        rows.append(
            {
                "pattern_label": pattern,
                "company_count": int(counts.get(pattern, 0)),
            }
        )
    # Append any patterns not in the canonical 8 (defensive — shouldn't happen)
    for label, cnt in counts.items():
        if label not in ALL_EIGHT_PATTERNS:
            rows.append({"pattern_label": label, "company_count": int(cnt)})
    return pd.DataFrame(rows)


# ---------------------------------------------------------------------------
# YoY pattern-change detection
# ---------------------------------------------------------------------------
def detect_pattern_changes(
    ca_df: pd.DataFrame,
    conn: sqlite3.Connection,
) -> pd.DataFrame:
    """Find companies whose capital-allocation pattern changed YoY in the latest
    transition (e.g. Reinvestor → Distress Signal between 2023-03 and 2024-03).

    Only reports the **most recent** transition per company (prev_year → latest_year).
    """
    # Get company name + sector
    meta_df = pd.read_sql(
        """
        SELECT c.id AS company_id, c.company_name, s.broad_sector AS sector
        FROM companies c
        LEFT JOIN sectors s ON s.company_id = c.id
        """,
        conn,
    )

    changes: list[PatternChangeRow] = []
    for cid, grp in ca_df.groupby("company_id", sort=False):
        grp = grp.sort_values("year").reset_index(drop=True)
        if len(grp) < 2:
            continue
        latest_row = grp.iloc[-1]
        prev_row = grp.iloc[-2]
        if latest_row["pattern_label"] != prev_row["pattern_label"]:
            meta = meta_df[meta_df["company_id"] == cid]
            company_name = meta["company_name"].iloc[0] if len(meta) else cid
            sector = meta["sector"].iloc[0] if len(meta) else None
            if isinstance(sector, float) and pd.isna(sector):
                sector = None
            changes.append(
                PatternChangeRow(
                    company_id=cid,
                    company_name=company_name,
                    sector=sector,
                    prev_year=str(prev_row["year"]),
                    prev_pattern=str(prev_row["pattern_label"]),
                    latest_year=str(latest_row["year"]),
                    latest_pattern=str(latest_row["pattern_label"]),
                )
            )

    out_df = pd.DataFrame([c.as_dict() for c in changes], columns=list(PATTERN_CHANGES_COLUMNS))
    logger.info("Detected %d YoY capital-allocation pattern changes", len(out_df))
    return out_df


# ---------------------------------------------------------------------------
# Writers
# ---------------------------------------------------------------------------
def write_pattern_changes_csv(df: pd.DataFrame, path: Path) -> int:
    """Write pattern_changes to CSV. Returns number of rows written."""
    path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(path, index=False)
    return len(df)


def write_pattern_distribution_xlsx(
    distribution: pd.DataFrame,
    changes_df: pd.DataFrame,
    completeness: dict,
    path: Path,
) -> Path:
    """Write a multi-sheet distribution summary workbook."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils.dataframe import dataframe_to_rows

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")

    # Sheet 1: Distribution
    ws1 = wb.active
    ws1.title = "Pattern Distribution"
    total = int(distribution["company_count"].sum())
    dist_with_pct = distribution.copy()
    dist_with_pct["pct_of_companies"] = dist_with_pct["company_count"].apply(
        lambda x: round(x / total * 100, 1) if total > 0 else 0.0
    )
    for r_idx, row in enumerate(dataframe_to_rows(dist_with_pct, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            cell = ws1.cell(row=r_idx, column=c_idx, value=value)
            if r_idx == 1:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")
    for col_cells in ws1.columns:
        max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        ws1.column_dimensions[col_cells[0].column_letter].width = min(max_len + 3, 40)
    ws1.freeze_panes = "A2"

    # Sheet 2: Pattern changes
    ws2 = wb.create_sheet("Pattern Changes (YoY)")
    for r_idx, row in enumerate(dataframe_to_rows(changes_df, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            cell = ws2.cell(row=r_idx, column=c_idx, value=value)
            if r_idx == 1:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")
    for col_cells in ws2.columns:
        max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        ws2.column_dimensions[col_cells[0].column_letter].width = min(max_len + 3, 30)
    ws2.freeze_panes = "A2"

    # Sheet 3: Completeness audit
    ws3 = wb.create_sheet("Completeness Audit")
    audit_rows = [
        ("Metric", "Value"),
        ("Total companies in DB", completeness["total_companies_db"]),
        ("Total companies in CSV", completeness["total_companies_csv"]),
        ("Total expected (company, year) rows", completeness["total_rows_expected"]),
        ("Total rows in CSV", completeness["total_rows_csv"]),
        ("Missing companies", ", ".join(completeness["missing_company_ids"]) or "(none)"),
        ("Missing rows count", len(completeness["missing_rows"])),
        ("Is complete?", "YES" if completeness["is_complete"] else "NO"),
    ]
    for r_idx, row in enumerate(audit_rows, 1):
        for c_idx, value in enumerate(row, 1):
            cell = ws3.cell(row=r_idx, column=c_idx, value=value)
            if r_idx == 1:
                cell.fill = header_fill
                cell.font = header_font
    ws3.column_dimensions["A"].width = 40
    ws3.column_dimensions["B"].width = 60

    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# High-level entry point
# ---------------------------------------------------------------------------
def run_capital_allocation_report(
    db_path: Path | str | None = None,
    output_dir: Path | str | None = None,
) -> tuple[dict, pd.DataFrame, pd.DataFrame, Path]:
    """Run the Day-32 Capital Allocation report.

    Returns ``(completeness_dict, distribution_df, changes_df, xlsx_path)``.
    """
    from src.analytics.cashflow_intelligence import run_cashflow_intelligence
    from src.utils.config import settings

    if db_path is None:
        db_path = settings.PROJECT_ROOT / "db" / "nifty100.db"
    if output_dir is None:
        output_dir = settings.PROJECT_ROOT / "output"
    db_path = Path(db_path)
    output_dir = Path(output_dir)

    ca_csv_path = output_dir / "capital_allocation.csv"
    if not ca_csv_path.exists():
        raise FileNotFoundError(
            f"capital_allocation.csv not found at {ca_csv_path} (run Day 11 KPI engine first)"
        )

    ca_df = pd.read_csv(ca_csv_path)
    conn = sqlite3.connect(str(db_path))
    try:
        completeness = verify_capital_allocation_completeness(ca_df, conn)
        distribution = build_pattern_distribution(ca_df)
        changes_df = detect_pattern_changes(ca_df, conn)

        # Write outputs
        write_pattern_changes_csv(changes_df, output_dir / "pattern_changes.csv")
        xlsx_path = write_pattern_distribution_xlsx(
            distribution,
            changes_df,
            completeness,
            output_dir / "capital_allocation_report.xlsx",
        )

        # Refresh cashflow_intelligence.xlsx to ensure capital_allocation_label is present
        run_cashflow_intelligence(db_path=db_path, output_dir=output_dir)
    finally:
        conn.close()

    return completeness, distribution, changes_df, xlsx_path


__all__ = [
    "ALL_EIGHT_PATTERNS",
    "PATTERN_CASH_ACCUMULATOR",
    "PATTERN_CHANGES_COLUMNS",
    "PATTERN_DISTRESS_SIGNAL",
    "PATTERN_GROWTH_FUNDED_BY_DEBT",
    "PATTERN_LIQUIDATING_ASSETS",
    "PATTERN_MIXED",
    "PATTERN_PRE_REVENUE",
    "PATTERN_REINVESTOR",
    "PATTERN_SHAREHOLDER_RETURNS",
    "PatternChangeRow",
    "build_pattern_distribution",
    "detect_pattern_changes",
    "run_capital_allocation_report",
    "verify_capital_allocation_completeness",
    "write_pattern_changes_csv",
    "write_pattern_distribution_xlsx",
]
